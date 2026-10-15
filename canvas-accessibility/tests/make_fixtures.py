#!/usr/bin/env python3
"""Generate synthetic test files for the accessibility checkers.

Every file is built from scratch with made-up content (no real course
material), so the folder is safe to publish. Each "bad" file breaks a known
set of rules; each "good" file should pass its checker cleanly.

Usage:  python tests/make_fixtures.py [output_dir]   (default: tests/fixtures)
"""
import io
import sys
from pathlib import Path

import pikepdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image
from pptx import Presentation
from pptx.util import Inches as PInches
from pptx.oxml.ns import qn as pqn


def png_bytes():
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), (70, 130, 180)).save(buf, "PNG")
    buf.seek(0)
    return buf


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def set_docdefault_lang(doc, value):
    """Set (value) or remove (None) the document-wide default language."""
    rpr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    for el in rpr.findall(qn("w:lang")):
        rpr.remove(el)
    if value:
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), value)
        rpr.append(lang)


# ---------------------------------------------------------------- docx
def docx_bad(path):
    doc = Document()
    for i in range(6):  # looks like headings + body, but no Heading styles
        p = doc.add_paragraph()
        p.add_run(f"Section {i + 1}" if i % 2 == 0 else "Body text for the section.").bold = (i % 2 == 0)
    t = doc.add_table(rows=3, cols=2)  # no header flag, no caption
    for r in range(3):
        for c in range(2):
            t.cell(r, c).text = f"r{r}c{c}"
    doc.add_picture(png_bytes(), width=Inches(1))  # no alt text
    pic2 = doc.add_picture(png_bytes(), width=Inches(1))
    pic2._inline.docPr.set("descr", "image2.png")  # filename used as alt text
    add_hyperlink(doc.add_paragraph("See the syllabus "), "https://example.edu/syllabus", "click here.")
    # Language only on one style, not in docDefaults (old check passed this)
    set_docdefault_lang(doc, None)
    normal_rpr = doc.styles["Normal"].element.get_or_add_rPr()
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    normal_rpr.append(lang)
    doc.core_properties.title = ""
    doc.save(path)


def docx_skip(path):
    doc = Document()
    doc.add_paragraph("Week 3 Worksheet")  # untagged title
    doc.add_heading("Activity 1", level=2)  # structure starts at Heading 2
    doc.add_paragraph("Do the activity.")
    doc.add_heading("Details", level=4)  # skips Heading 3
    doc.core_properties.title = "Week 3 Worksheet"
    doc.save(path)


def docx_good(path):
    doc = Document()
    doc.add_heading("Sample Handout", level=0)
    doc.add_heading("Overview", level=1)
    doc.add_paragraph("Body text.")
    doc.add_heading("Details", level=2)
    t = doc.add_table(rows=3, cols=2)
    for r in range(3):
        for c in range(2):
            t.cell(r, c).text = ["Term", "Meaning"][c] if r == 0 else f"value {r}{c}"
    trpr = t.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:tblHeader"))
    tblpr = t._tbl.tblPr
    cap = OxmlElement("w:tblCaption"); cap.set(qn("w:val"), "Glossary")
    desc = OxmlElement("w:tblDescription"); desc.set(qn("w:val"), "Two-column table of terms and meanings.")
    tblpr.append(cap); tblpr.append(desc)
    doc.add_picture(png_bytes(), width=Inches(1))._inline.docPr.set("descr", "Blue rectangle used as a sample figure")
    add_hyperlink(doc.add_paragraph("Read "), "https://example.edu/syllabus", "the course syllabus (PDF)")
    set_docdefault_lang(doc, "en-US")
    doc.core_properties.title = "Sample Handout"
    doc.save(path)


# ---------------------------------------------------------------- pptx
def pptx_bad(path):
    prs = Presentation()
    blank, title_only = prs.slide_layouts[6], prs.slide_layouts[5]

    s1 = prs.slides.add_slide(blank)  # no title placeholder
    s1.shapes.add_textbox(PInches(1), PInches(0.5), PInches(6), PInches(1)).text = "Looks like a title"
    pic = s1.shapes.add_picture(png_bytes(), PInches(1), PInches(2))
    pic._element.find(".//" + pqn("p:cNvPr")).set("descr", "IMG_4521.png")  # filename alt text

    s2 = prs.slides.add_slide(title_only)
    s2.shapes.title.text = "Overview"
    grp = s2.shapes.add_group_shape()
    gpic = grp.shapes.add_picture(png_bytes(), PInches(1), PInches(2))  # grouped, no alt text
    gpic._element.find(".//" + pqn("p:cNvPr")).set("descr", "")
    tbl = s2.shapes.add_table(3, 2, PInches(4), PInches(2), PInches(4), PInches(1.5)).table
    tbl._tbl.tblPr.set("firstRow", "0")

    s3 = prs.slides.add_slide(title_only)
    s3.shapes.title.text = "Overview"  # duplicate title
    prs.save(path)


def pptx_good(path):
    prs = Presentation()
    layout = prs.slide_layouts[5]
    for name in ("Introduction", "Results"):
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = name
    pic = prs.slides[0].shapes.add_picture(png_bytes(), PInches(1), PInches(2))
    pic._element.find(".//" + pqn("p:cNvPr")).set("descr", "Bar chart of weekly quiz averages")
    tbl = prs.slides[1].shapes.add_table(3, 2, PInches(1), PInches(2), PInches(4), PInches(1.5)).table
    tbl._tbl.tblPr.set("firstRow", "1")
    prs.save(path)


# ---------------------------------------------------------------- pdf
def pdf_bad(path):
    pdf = pikepdf.new()
    pdf.add_blank_page()
    pdf.docinfo["/Title"] = "Microsoft Word - handout.docx"
    pdf.save(path)


def pdf_good(path):
    pdf = pikepdf.new()
    pdf.add_blank_page()
    pdf.Root.MarkInfo = pikepdf.Dictionary(Marked=True)
    pdf.Root.StructTreeRoot = pdf.make_indirect(pikepdf.Dictionary(Type=pikepdf.Name.StructTreeRoot))
    pdf.Root.Lang = pikepdf.String("en-US")
    pdf.docinfo["/Title"] = "Sample Handout"
    pdf.save(path)


BUILDERS = {
    "docx_bad.docx": docx_bad, "docx_skip.docx": docx_skip, "docx_good.docx": docx_good,
    "pptx_bad.pptx": pptx_bad, "pptx_good.pptx": pptx_good,
    "pdf_bad.pdf": pdf_bad, "pdf_good.pdf": pdf_good,
}

if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).parent / "fixtures")
    out.mkdir(parents=True, exist_ok=True)
    for name, fn in BUILDERS.items():
        fn(out / name)
        print("wrote", out / name)
