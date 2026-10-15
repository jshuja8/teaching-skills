#!/usr/bin/env python3
"""
Read-only accessibility checker for .docx files.

Checks for the issues Canvas's accessibility checker and screen readers care
about most: heading structure, table header rows, table/image alt text,
non-descriptive hyperlink text, document title, and default language.

Does NOT modify the file. Usage:
    python check_docx_accessibility.py <file.docx>

Prints a plain-text report and exits 1 if any issues were found (0 if clean),
so it can be used as a quick pass/fail gate as well as a report.
"""
import re
import string
import sys
from lxml import etree
from docx import Document
from docx.oxml.ns import qn

GENERIC_LINK_TEXT = {
    "click here", "here", "link", "read more", "more", "this link", "click"
}

FILENAME_LIKE = re.compile(r"\.(png|jpe?g|gif|bmp|tiff?|emf|wmf|svg)$", re.IGNORECASE)
GENERIC_SHAPE_NAME = re.compile(r"^(picture|image|graphic|chart)\s*\d*$", re.IGNORECASE)


def normalize_link_text(text):
    """Lowercase and strip surrounding punctuation, so 'Click here.' matches."""
    return text.strip().strip(string.punctuation + " \u00a0").lower()


HEADING_LEVELS = {
    "Title": 0, "Heading 1": 1, "Heading 2": 2, "Heading 3": 3,
    "Heading 4": 4, "Heading 5": 5, "Heading 6": 6,
}


def check(path):
    doc = Document(path)
    issues = []   # list of (title, detail)
    passes = []   # list of strings

    # 1. Heading structure
    heading_styles_used = set()
    heading_sequence = []  # (style_name, text) in document order
    body_paragraphs_with_text = 0
    for p in doc.paragraphs:
        if p.text.strip():
            body_paragraphs_with_text += 1
        style_name = p.style.name if p.style is not None else ""
        if style_name == "Title" or style_name.startswith("Heading"):
            heading_styles_used.add(style_name)
            if style_name in HEADING_LEVELS:
                heading_sequence.append((style_name, p.text.strip()[:50]))

    if body_paragraphs_with_text > 5 and not heading_styles_used:
        issues.append((
            "No heading structure",
            "The document has substantial text but no paragraph uses a Heading or "
            "Title style. Bold/large text LOOKS like a heading but isn't tagged as "
            "one -- screen readers and Canvas's outline/navigation view rely on the "
            "actual style, not the visual formatting."
        ))
    elif heading_styles_used:
        passes.append(f"Heading styles in use: {', '.join(sorted(heading_styles_used))}")

    # 1b. Heading level skips -- including the easy-to-miss case where the
    # document's OWN title/subtitle is left untagged while sections below it
    # are correctly tagged (so "some headings exist" passes the check above,
    # but the very top of the structure still skips levels).
    if heading_sequence:
        first_level = HEADING_LEVELS[heading_sequence[0][0]]
        if first_level > 1:
            issues.append((
                f'Heading structure starts at "{heading_sequence[0][0]}" (skips Title/Heading 1)',
                f'The first tagged heading in the document is "{heading_sequence[0][1]}" '
                f'at level {heading_sequence[0][0]}, with nothing above it. This usually '
                "means the document's own title/subtitle line(s) were left as plain text "
                "while sections below were correctly tagged -- easy to miss because "
                "headings clearly exist elsewhere. Add Title (and Heading 1 for a "
                "subtitle, if there is one) to the top of the document."
            ))
        prev_level = first_level
        for style_name, text in heading_sequence[1:]:
            lvl = HEADING_LEVELS[style_name]
            if lvl > prev_level + 1:
                issues.append((
                    f'Heading level skipped before "{text}"',
                    f'This heading is tagged "{style_name}", but the previous heading was '
                    f'only level {prev_level} -- jumping more than one level breaks the '
                    "outline a screen reader user navigates by. Use the next level down "
                    "instead, or add the missing intermediate heading."
                ))
            prev_level = lvl

    # 2. Tables: header rows + alt text (caption/description)
    for i, table in enumerate(doc.tables, start=1):
        tbl = table._tbl
        rows = tbl.findall(qn("w:tr"))
        if not rows:
            continue

        first_row_trPr = rows[0].find(qn("w:trPr"))
        has_header_flag = (
            first_row_trPr is not None
            and first_row_trPr.find(qn("w:tblHeader")) is not None
        )
        n_rows = len(rows)
        n_cells_first_row = len(rows[0].findall(qn("w:tc")))
        looks_like_data_table = n_rows > 1 and n_cells_first_row > 1

        if looks_like_data_table and not has_header_flag:
            issues.append((
                f"Table {i}: first row may need header marking -- verify before fixing",
                "This table has multiple rows and columns, which is necessary but NOT "
                "sufficient for needing a header row. Read the actual cell content: "
                "mark the first row as a header (w:tblHeader) ONLY if it functions as "
                "a column label that's still true for every row below it (e.g. "
                "'Kitchen concept | Computer concept' above a list of matched pairs). "
                "If the first row is itself a data value specific to that row (e.g. a "
                "positional label like 'Top' paired with content, where later rows have "
                "a blank or different value in that same column), it is NOT a header "
                "row -- marking it as one makes a screen reader misdescribe every row "
                "below it. When in doubt, give the table alt text instead and leave "
                "the header flag off."
            ))
        elif looks_like_data_table:
            passes.append(f"Table {i}: header row correctly marked")

        tblPr = tbl.find(qn("w:tblPr"))
        has_caption = tblPr is not None and tblPr.find(qn("w:tblCaption")) is not None
        if not has_caption:
            issues.append((
                f"Table {i}: no alt text",
                "Give this table a short title + description (Table Properties > "
                "Alt Text in Word, stored as tblCaption/tblDescription) so a screen "
                "reader announces what it's for before reading the cells."
            ))

    # 3. Images: alt text (skip anything marked decorative)
    body = doc.element.body
    for docPr in body.iter(qn("wp:docPr")):
        name = docPr.get("name") or "image"
        descr = (docPr.get("descr") or "").strip()
        title_attr = (docPr.get("title") or "").strip()
        subtree_xml = etree.tostring(docPr, encoding="unicode").lower()
        marked_decorative = "decorative" in subtree_xml
        if not descr and not title_attr and not marked_decorative:
            issues.append((
                f"Image \"{name}\": no alt text",
                "This image has no alternative text and isn't marked decorative. "
                "If it conveys information (a diagram, screenshot, chart), add a "
                "concise description. If it's purely visual flourish, mark it "
                "decorative instead of leaving it blank."
            ))
        elif (descr or title_attr) and not marked_decorative:
            alt_text = descr or title_attr
            if FILENAME_LIKE.search(alt_text) or GENERIC_SHAPE_NAME.match(alt_text):
                issues.append((
                    f'Image "{name}": alt text is just a filename/placeholder ("{alt_text}")',
                    "Word sometimes fills alt text with the source filename or a "
                    "generic label. That passes a blank check but tells a screen "
                    "reader user nothing. Replace it with a real description."
                ))

    # 4. Hyperlinks with non-descriptive visible text
    for hyperlink in body.iter(qn("w:hyperlink")):
        text = "".join(t.text or "" for t in hyperlink.iter(qn("w:t"))).strip()
        if normalize_link_text(text) in GENERIC_LINK_TEXT:
            issues.append((
                f"Hyperlink text \"{text}\"",
                "Screen reader users often scan a list of links out of context. "
                "Generic text like this doesn't say where the link goes -- use "
                "something descriptive, e.g. 'Week 1 syllabus (PDF)' instead of "
                "'click here'."
            ))

    # 5. Document title (core properties)
    core_title = (doc.core_properties.title or "").strip()
    if not core_title:
        issues.append((
            "No document title set",
            "The file's Title property (File > Info in Word) is empty, so Canvas "
            "and screen readers may announce the raw filename instead of a "
            "readable title."
        ))
    else:
        passes.append(f'Document title set: "{core_title}"')

    # 6. Default language -- must be in docDefaults (what assistive technology
    # uses document-wide), not merely on one style somewhere.
    def has_lang(el):
        return el is not None and any(l.get(qn("w:val")) for l in el.iter(qn("w:lang")))

    doc_defaults = doc.styles.element.find(qn("w:docDefaults"))
    if has_lang(doc_defaults):
        passes.append("Default document language set (docDefaults)")
    elif has_lang(doc.styles.element):
        issues.append((
            "Language set only on individual styles, not as the document default",
            "Some styles carry a language tag, but the document-wide default "
            "(w:docDefaults > w:rPrDefault in styles.xml) does not, so text in "
            "other styles may be read with the wrong pronunciation rules. Set the "
            "default language (Review > Language > Set Proofing Language, with "
            "'Set As Default')."
        ))
    else:
        issues.append((
            "No default language set",
            "Without a language tag, screen readers may apply the wrong "
            "pronunciation rules to the text."
        ))

    return issues, passes


def main():
    if len(sys.argv) != 2:
        print("Usage: python check_docx_accessibility.py <file.docx>")
        sys.exit(2)

    path = sys.argv[1]
    issues, passes = check(path)

    print(f"=== DOCX accessibility check: {path} ===\n")
    print(f"{len(issues)} issue(s) found, {len(passes)} check(s) passed.\n")

    if issues:
        print("ISSUES")
        print("-" * 60)
        for title, detail in issues:
            print(f"[FAIL] {title}")
            print(f"       {detail}\n")

    if passes:
        print("PASSED")
        print("-" * 60)
        for p in passes:
            print(f"[OK]   {p}")

    sys.exit(1 if issues else 0)


if __name__ == "__main__":
    main()
